"""
Views for the billing application.
"""
import io
import json
import logging
import os
from decimal import Decimal
from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db import transaction
from django.db.models import Q, Count, Sum, F, DecimalField, Value
from django.db.models.functions import Coalesce
from django.http import HttpResponse, HttpResponseRedirect
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse, reverse_lazy
from django.utils import timezone
from django.views import View
from django.views.generic import ListView, DetailView, DeleteView
from xhtml2pdf import pisa
from .forms import (
    BillForm,
    BillItemFormSet,
    MarginMoneyDetailForm,
    DeliveryChallanDetailForm,
    ProductForm,
    TractorStockUnitForm,
    PurchaseBillForm,
    PurchaseBillItemFormSet,
)
from .models import (
    Bill, BillItem, MarginMoneyDetail, DeliveryChallanDetail,
    Product, TractorStockUnit, PurchaseBill, PurchaseBillItem
)
from .inventory_service import (
    deduct_stock_for_bill,
    restore_stock_for_bill,
    apply_purchase_stock_inward,
    validate_stock_availability,
    find_matching_product,
)
from customers.models import Customer
from settings_app.models import CompanyProfile, GSTConfig, NotificationConfig
from settings_app.notification_service import (
    normalize_indian_mobile,
    build_invoice_alert_message,
    send_notification,
)
from .utils import link_callback

logger = logging.getLogger(__name__)


def get_billing_meta():
    """Returns serialized GST configurations, customer state mappings, and in-stock tractors for live Alpine.js math."""
    company = CompanyProfile.get_instance()
    gst_configs = {}
    for g in GSTConfig.objects.filter(is_active=True):
        gst_configs[str(g.id)] = {
            'name': g.name,
            'hsn': g.hsn_code or '8701',
            'cgst': float(g.cgst_rate or 0),
            'sgst': float(g.sgst_rate or 0),
            'igst': float(g.igst_rate or (g.cgst_rate + g.sgst_rate)),
        }
    customer_states = {}
    for c in Customer.objects.filter(is_active=True):
        customer_states[str(c.id)] = (c.state or '').strip().lower()

    in_stock_tractors = list(TractorStockUnit.objects.filter(status='IN_STOCK').select_related('product'))
    available_tractors = [
        {
            'id': u.id,
            'chassis': u.chassis_number,
            'engine': u.engine_number,
            'model': u.product.name,
            'price': float(u.product.default_price or 0),
            'hsn': u.product.hsn_code or '8701',
        }
        for u in in_stock_tractors
    ]

    catalog_items = []
    # 1. Serialized Physical Units in Showroom Inventory
    for u in in_stock_tractors:
        catalog_items.append({
            'name': f"{u.product.name} (Chassis: {u.chassis_number})",
            'model': u.product.name,
            'category': 'Serialized Tractor',
            'category_code': 'STOCK',
            'hsn': u.product.hsn_code or '8701',
            'price': float(u.product.default_price or 0),
            'chassis': u.chassis_number,
            'engine': u.engine_number,
            'is_stock': True,
            'stock_qty': 1,
            'is_in_stock': True,
            'product_id': u.product_id,
            'unit_id': u.id,
            'part_number': '',
        })
    # 2. All catalog products (Tractors, Implements, Trolleys, Spares, Lubricants, Accessories)
    for p in Product.objects.filter(is_active=True).order_by('category', 'name'):
        catalog_items.append({
            'name': p.name,
            'model': p.name,
            'category': p.get_category_display(),
            'category_code': p.category,
            'hsn': p.hsn_code or ('8701' if p.category == 'TRACTOR' else '8432'),
            'price': float(p.default_price or 0),
            'chassis': '',
            'engine': '',
            'is_stock': False,
            'stock_qty': p.stock_quantity,
            'is_in_stock': p.stock_quantity > 0,
            'product_id': p.id,
            'unit_id': None,
            'part_number': p.part_number or '',
        })

    return {
        'gst_configs_json': json.dumps(gst_configs),
        'customer_states_json': json.dumps(customer_states),
        'company_state': (company.state or 'Maharashtra').strip().lower(),
        'available_tractors_json': json.dumps(available_tractors),
        'catalog_items_json': json.dumps(catalog_items),
    }


class BillListView(LoginRequiredMixin, ListView):
    """
    Paginated list of all bills with filters for type, status, and search query.
    """
    model = Bill
    template_name = 'billing/bill_list.html'
    context_object_name = 'bills'
    paginate_by = 15

    def get_paginate_by(self, queryset):
        page_size = self.request.GET.get('page_size')
        if page_size and page_size.isdigit():
            size = int(page_size)
            if 5 <= size <= 100:
                return size
        return self.paginate_by

    def get_queryset(self):
        qs = Bill.objects.select_related('customer').all().order_by('-created_at')
        q = self.request.GET.get('q', '').strip()
        bill_type = self.request.GET.get('type', '').strip()
        status = self.request.GET.get('status', '').strip()
        payment_status = self.request.GET.get('payment_status', '').strip()

        if q:
            qs = qs.filter(
                Q(bill_number__icontains=q) |
                Q(customer__full_name__icontains=q) |
                Q(customer__mobile_primary__icontains=q) |
                Q(customer__village__icontains=q)
            )
        if bill_type:
            qs = qs.filter(bill_type=bill_type)
        if status:
            qs = qs.filter(status=status)
        if payment_status:
            qs = qs.annotate(
                annotated_paid=Coalesce(
                    Sum('payment_allocations__allocated_amount'),
                    Value(Decimal('0.00')),
                    output_field=DecimalField()
                )
            )
            if payment_status == 'PAID':
                qs = qs.filter(Q(annotated_paid__gte=F('grand_total')) | Q(grand_total__lte=0))
            elif payment_status == 'PARTIAL':
                qs = qs.filter(annotated_paid__gt=0, annotated_paid__lt=F('grand_total'), grand_total__gt=0)
            elif payment_status == 'UNPAID':
                qs = qs.filter(annotated_paid__lte=0, grand_total__gt=0)

        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['current_page_size'] = self.get_paginate_by(None)
        context['current_type'] = self.request.GET.get('type', '')
        context['current_status'] = self.request.GET.get('status', '')
        context['current_payment_status'] = self.request.GET.get('payment_status', '')
        context['search_query'] = self.request.GET.get('q', '')
        context['bill_types'] = Bill.BILL_TYPE_CHOICES

        # KPI Overview Metrics using fast DB aggregation
        kpis = Bill.objects.aggregate(
            total_bills_count=Count('id'),
            tax_invoices_count=Count('id', filter=Q(bill_type='TAX_INVOICE')),
            tax_invoices_sum=Coalesce(Sum('grand_total', filter=Q(bill_type='TAX_INVOICE')), Value(Decimal('0.00')), output_field=DecimalField()),
            quotations_count=Count('id', filter=Q(bill_type='QUOTATION')),
            quotations_sum=Coalesce(Sum('grand_total', filter=Q(bill_type='QUOTATION')), Value(Decimal('0.00')), output_field=DecimalField()),
            margin_money_count=Count('id', filter=Q(bill_type='MARGIN_MONEY')),
            margin_money_sum=Coalesce(Sum('grand_total', filter=Q(bill_type='MARGIN_MONEY')), Value(Decimal('0.00')), output_field=DecimalField()),
            delivery_challan_count=Count('id', filter=Q(bill_type='DELIVERY_CHALLAN')),
        )

        unpaid_invoices_qs = (
            Bill.objects.filter(bill_type='TAX_INVOICE')
            .exclude(status='CANCELLED')
            .annotate(
                annotated_paid=Coalesce(
                    Sum('payment_allocations__allocated_amount'),
                    Value(Decimal('0.00')),
                    output_field=DecimalField()
                )
            )
            .filter(annotated_paid__lt=F('grand_total'), grand_total__gt=0)
        )
        unpaid_stats = unpaid_invoices_qs.aggregate(
            unpaid_count=Count('id'),
            unpaid_amount=Coalesce(Sum(F('grand_total') - F('annotated_paid')), Value(Decimal('0.00')), output_field=DecimalField()),
        )

        context['total_bills_count'] = kpis['total_bills_count']
        context['total_billed_sum'] = kpis['tax_invoices_sum']
        context['tax_invoices_count'] = kpis['tax_invoices_count']
        context['tax_invoices_sum'] = kpis['tax_invoices_sum']
        context['quotations_count'] = kpis['quotations_count']
        context['quotations_sum'] = kpis['quotations_sum']
        context['margin_money_count'] = kpis['margin_money_count']
        context['margin_money_sum'] = kpis['margin_money_sum']
        context['delivery_challan_count'] = kpis['delivery_challan_count']
        context['unpaid_invoices_count'] = unpaid_stats['unpaid_count']
        context['unpaid_invoices_amount'] = unpaid_stats['unpaid_amount']

        return context


class BillCreateView(LoginRequiredMixin, View):
    """
    Unified bill creation view handling Bill, BillItems formset,
    and type-specific details (MarginMoney or DeliveryChallan).
    """
    template_name = 'billing/bill_form.html'

    def get_initial_type(self):
        t = self.request.GET.get('type', 'QUOTATION').upper()
        valid = [k for k, v in Bill.BILL_TYPE_CHOICES]
        return t if t in valid else 'QUOTATION'

    def get(self, request, *args, **kwargs):
        initial_type = self.get_initial_type()
        customer_id = request.GET.get('customer')
        product_id = request.GET.get('product_id')
        unit_id = request.GET.get('unit_id')

        initial = {'bill_type': initial_type}
        if customer_id:
            try:
                initial['customer'] = Customer.objects.get(pk=customer_id)
            except Customer.DoesNotExist:
                pass

        form = BillForm(initial=initial)

        # Support instant billing from inventory
        initial_items = []
        if unit_id:
            try:
                unit = TractorStockUnit.objects.select_related('product').get(pk=unit_id)
                initial_items.append({
                    'item_description': unit.product.name,
                    'chassis_number': unit.chassis_number,
                    'engine_number': unit.engine_number,
                    'hsn_code': unit.product.hsn_code or '8701',
                    'quantity': 1,
                    'unit_price': unit.product.default_price,
                    'discount': Decimal('0.00'),
                })
            except TractorStockUnit.DoesNotExist:
                pass
        elif product_id:
            try:
                prod = Product.objects.get(pk=product_id)
                initial_items.append({
                    'item_description': prod.name,
                    'hsn_code': prod.hsn_code or ('8701' if prod.category == 'TRACTOR' else '8432'),
                    'quantity': 1,
                    'unit_price': prod.default_price,
                    'discount': Decimal('0.00'),
                })
            except Product.DoesNotExist:
                pass

        if initial_items:
            formset = BillItemFormSet(initial=initial_items)
        else:
            formset = BillItemFormSet()

        margin_form = MarginMoneyDetailForm()
        challan_form = DeliveryChallanDetailForm()

        context = {
            'form': form,
            'formset': formset,
            'margin_form': margin_form,
            'challan_form': challan_form,
            'bill_type': initial_type,
            'is_edit': False,
            'page_title': f"Generate {dict(Bill.BILL_TYPE_CHOICES).get(initial_type, 'Bill')}",
        }
        context.update(get_billing_meta())
        return render(request, self.template_name, context)

    def post(self, request, *args, **kwargs):
        form = BillForm(request.POST)
        formset = BillItemFormSet(request.POST)
        margin_form = MarginMoneyDetailForm(request.POST)
        challan_form = DeliveryChallanDetailForm(request.POST)

        if form.is_valid() and formset.is_valid():
            bill_type = form.cleaned_data.get('bill_type')
            valid_extra = True

            if bill_type == 'MARGIN_MONEY':
                valid_extra = margin_form.is_valid()
            elif bill_type == 'DELIVERY_CHALLAN':
                valid_extra = challan_form.is_valid()

            if valid_extra:
                with transaction.atomic():
                    bill = form.save(commit=False)
                    bill.created_by = request.user
                    bill.save()

                    # Save items
                    formset.instance = bill
                    items = formset.save(commit=False)
                    for i, item in enumerate(items, start=1):
                        item.sr_no = i
                        item.save()
                    formset.save_m2m()

                    # Save extra detail if applicable
                    if bill_type == 'MARGIN_MONEY':
                        m = margin_form.save(commit=False)
                        m.bill = bill
                        m.save()
                        # Link and mark tractor booked if available in stock
                        if m.tractor_model:
                            matched_unit = TractorStockUnit.objects.filter(
                                product__name__icontains=m.tractor_model.strip(),
                                status='IN_STOCK'
                            ).first()
                            if matched_unit:
                                matched_unit.allocated_bill = bill
                                matched_unit.status = 'BOOKED'
                                matched_unit.save(update_fields=['allocated_bill', 'status'])
                                if matched_unit.product:
                                    matched_unit.product.stock_quantity = TractorStockUnit.objects.filter(product=matched_unit.product, status='IN_STOCK').count()
                                    matched_unit.product.save(update_fields=['stock_quantity'])

                    elif bill_type == 'DELIVERY_CHALLAN':
                        c = challan_form.save(commit=False)
                        c.bill = bill
                        c.save()

                    # Deduct stock for commercial documents (Tax Invoice & Delivery Challan)
                    deduct_stock_for_bill(bill)

                    # Calculate totals and words
                    bill.calculate_totals(save=True)

                # Auto-send WhatsApp/SMS document alert if enabled
                config = NotificationConfig.get_instance()
                if config.is_enabled and config.auto_send_invoice_alert and bill.status != 'CANCELLED':
                    try:
                        recipient, message = build_invoice_alert_message(bill)
                        log_entry = send_notification(
                            recipient=recipient,
                            message=message,
                            notification_type='INVOICE_ALERT',
                            user=request.user,
                            entity_id=bill.bill_number
                        )
                        if log_entry and log_entry.status in ('SENT', 'SIMULATED'):
                            messages.info(request, f"WhatsApp / SMS notification sent to {recipient}.")
                    except Exception as e:
                        logger.warning(f"Could not auto-send bill notification: {e}")

                messages.success(request, f"{bill.get_bill_type_display()} '{bill.bill_number}' created successfully! Showroom stock updated.")
                return redirect('billing:bill_detail', pk=bill.pk)

        messages.error(request, "Please review and correct the highlighted errors below before generating document.")
        bill_type = request.POST.get('bill_type', 'QUOTATION')
        context = {
            'form': form,
            'formset': formset,
            'margin_form': margin_form,
            'challan_form': challan_form,
            'bill_type': bill_type,
            'is_edit': False,
            'page_title': f"Generate {dict(Bill.BILL_TYPE_CHOICES).get(bill_type, 'Bill')}",
        }
        context.update(get_billing_meta())
        return render(request, self.template_name, context)


class BillUpdateView(LoginRequiredMixin, View):
    """
    Update existing bill.
    """
    template_name = 'billing/bill_form.html'

    def get(self, request, pk, *args, **kwargs):
        bill = get_object_or_404(Bill, pk=pk)
        form = BillForm(instance=bill)
        formset = BillItemFormSet(instance=bill)

        margin_form = None
        if hasattr(bill, 'margin_detail'):
            margin_form = MarginMoneyDetailForm(instance=bill.margin_detail)
        elif bill.bill_type == 'MARGIN_MONEY':
            margin_form = MarginMoneyDetailForm()

        challan_form = None
        if hasattr(bill, 'challan_detail'):
            challan_form = DeliveryChallanDetailForm(instance=bill.challan_detail)
        elif bill.bill_type == 'DELIVERY_CHALLAN':
            challan_form = DeliveryChallanDetailForm()

        context = {
            'form': form,
            'formset': formset,
            'margin_form': margin_form,
            'challan_form': challan_form,
            'bill': bill,
            'bill_type': bill.bill_type,
            'is_edit': True,
            'page_title': f"Edit {bill.get_bill_type_display()} - {bill.bill_number}",
        }
        context.update(get_billing_meta())
        return render(request, self.template_name, context)

    def post(self, request, pk, *args, **kwargs):
        bill = get_object_or_404(Bill, pk=pk)
        form = BillForm(request.POST, instance=bill)
        formset = BillItemFormSet(request.POST, instance=bill)

        target_bill_type = request.POST.get('bill_type') or bill.bill_type
        margin_instance = getattr(bill, 'margin_detail', None)
        margin_form = MarginMoneyDetailForm(request.POST, instance=margin_instance) if target_bill_type == 'MARGIN_MONEY' else None

        challan_instance = getattr(bill, 'challan_detail', None)
        challan_form = DeliveryChallanDetailForm(request.POST, instance=challan_instance) if target_bill_type == 'DELIVERY_CHALLAN' else None

        if form.is_valid() and formset.is_valid():
            valid_extra = True
            if margin_form and not margin_form.is_valid():
                valid_extra = False
            if challan_form and not challan_form.is_valid():
                valid_extra = False

            if valid_extra:
                with transaction.atomic():
                    # Restore previous stock allocations before applying updated items
                    restore_stock_for_bill(bill)

                    bill = form.save()
                    formset.save()

                    if margin_form:
                        m = margin_form.save(commit=False)
                        m.bill = bill
                        m.save()
                    if challan_form:
                        c = challan_form.save(commit=False)
                        c.bill = bill
                        c.save()

                    # Deduct showroom stock for updated commercial document
                    deduct_stock_for_bill(bill)

                    if bill.bill_type == 'MARGIN_MONEY':
                        for item in bill.items.all():
                            if item.chassis_number:
                                unit = TractorStockUnit.objects.filter(chassis_number__iexact=item.chassis_number.strip()).first()
                                if unit and unit.status == 'IN_STOCK':
                                    unit.status = 'BOOKED'
                                    unit.allocated_bill = bill
                                    unit.save(update_fields=['status', 'allocated_bill'])

                    bill.calculate_totals(save=True)

                messages.success(request, f"Bill '{bill.bill_number}' updated successfully! Showroom stock synchronized.")
                return redirect('billing:bill_detail', pk=bill.pk)

        messages.error(request, "Please review and correct the highlighted errors below before saving.")
        context = {
            'form': form,
            'formset': formset,
            'margin_form': margin_form,
            'challan_form': challan_form,
            'bill': bill,
            'bill_type': bill.bill_type,
            'is_edit': True,
            'page_title': f"Edit {bill.get_bill_type_display()} - {bill.bill_number}",
        }
        context.update(get_billing_meta())
        return render(request, self.template_name, context)


class BillDetailView(LoginRequiredMixin, DetailView):
    """
    Displays the bill matching the official showroom quotation/invoice layout.
    """
    model = Bill
    template_name = 'billing/bill_detail.html'
    context_object_name = 'bill'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        bill = self.object
        items = list(bill.items.all().order_by('sr_no'))
        context['company'] = CompanyProfile.get_instance()
        context['items'] = items
        context['total_quantity'] = sum((item.quantity for item in items), 0)
        context['hsn_summary'] = bill.get_hsn_summary()
        margin_suggested = (bill.grand_total * Decimal('0.15')).quantize(Decimal('0.01'))
        context['margin_suggested'] = margin_suggested
        context['loan_balance'] = (bill.grand_total - margin_suggested).quantize(Decimal('0.01'))
        return context


class BillSendNotificationView(LoginRequiredMixin, View):
    """
    Manual trigger to dispatch or re-send bill / invoice document via WhatsApp / SMS.
    """
    def post(self, request, pk, *args, **kwargs):
        bill = get_object_or_404(Bill, pk=pk)
        try:
            recipient, message = build_invoice_alert_message(bill)
            log_entry = send_notification(
                recipient=recipient,
                message=message,
                notification_type='INVOICE_ALERT',
                user=request.user,
                entity_id=bill.bill_number
            )
            if log_entry.status == 'SENT':
                messages.success(request, f"Document alert sent to {recipient} via WhatsApp/SMS successfully!")
            elif log_entry.status == 'SIMULATED':
                messages.info(request, f"Document alert simulated for {recipient} (Simulator mode active).")
            else:
                messages.warning(request, f"Delivery warning: {log_entry.error_message}")
        except ValueError as ve:
            messages.error(request, str(ve))
        except Exception as e:
            logger.error(f"Error dispatching bill notification: {e}", exc_info=True)
            messages.error(request, f"Failed to send notification: {e}")

        return redirect('billing:bill_detail', pk=bill.pk)


class BillDeleteView(LoginRequiredMixin, DeleteView):
    """Delete a bill."""
    model = Bill
    template_name = 'billing/bill_confirm_delete.html'
    success_url = reverse_lazy('billing:bill_list')

    def form_valid(self, form):
        bill = self.get_object()
        restore_stock_for_bill(bill)
        messages.success(self.request, f"Bill '{bill.bill_number}' was deleted and allocated showroom stock was restored.")
        return super().form_valid(form)


@login_required
def bill_duplicate(request, pk):
    """
    Duplicates an existing bill with all its items and serial details as a new draft.
    """
    original = get_object_or_404(Bill, pk=pk)
    with transaction.atomic():
        new_bill = Bill.objects.create(
            bill_type=original.bill_type,
            customer=original.customer,
            bill_date=timezone.localdate(),
            due_date=original.due_date,
            supplier_ref=original.supplier_ref,
            other_ref=original.other_ref,
            bank_name=original.bank_name,
            accessories_insurance=original.accessories_insurance,
            is_gst_bill=original.is_gst_bill,
            gst_config=original.gst_config,
            discount_percentage=original.discount_percentage,
            discount_amount=original.discount_amount,
            notes=f"Copy of {original.bill_number}. {original.notes or ''}".strip(),
            status='DRAFT',
            created_by=request.user,
        )
        for item in original.items.all():
            BillItem.objects.create(
                bill=new_bill,
                sr_no=item.sr_no,
                product=item.product,
                tractor_unit=item.tractor_unit,
                item_description=item.item_description,
                chassis_number=item.chassis_number,
                engine_number=item.engine_number,
                hsn_code=item.hsn_code,
                quantity=item.quantity,
                unit_price=item.unit_price,
                discount=item.discount,
            )
        new_bill.calculate_totals(save=True)

    messages.success(request, f"Created duplicate bill '{new_bill.bill_number}' from '{original.bill_number}'.")
    return redirect('billing:bill_edit', pk=new_bill.pk)


@login_required
def bill_convert(request, pk):
    """
    Converts a Quotation into an official Tax Invoice copying all line items and equipment specs.
    """
    quotation = get_object_or_404(Bill, pk=pk)
    if quotation.bill_type != 'QUOTATION':
        messages.warning(request, "Only Quotations can be converted to Tax Invoices.")
        return redirect('billing:bill_detail', pk=quotation.pk)

    with transaction.atomic():
        invoice = Bill.objects.create(
            bill_type='TAX_INVOICE',
            customer=quotation.customer,
            bill_date=timezone.localdate(),
            supplier_ref=quotation.supplier_ref,
            other_ref=quotation.other_ref,
            bank_name=quotation.bank_name,
            accessories_insurance=quotation.accessories_insurance,
            is_gst_bill=quotation.is_gst_bill,
            gst_config=quotation.gst_config if quotation.is_gst_bill else None,
            discount_percentage=quotation.discount_percentage,
            discount_amount=quotation.discount_amount,
            notes=f"Converted from Quotation {quotation.bill_number}. {quotation.notes or ''}".strip(),
            status='DRAFT',
            created_by=request.user,
        )
        for item in quotation.items.all():
            BillItem.objects.create(
                bill=invoice,
                sr_no=item.sr_no,
                product=item.product,
                tractor_unit=item.tractor_unit,
                item_description=item.item_description,
                chassis_number=item.chassis_number,
                engine_number=item.engine_number,
                hsn_code=item.hsn_code or '8701',
                quantity=item.quantity,
                unit_price=item.unit_price,
                discount=item.discount,
            )
        invoice.calculate_totals(save=True)

    messages.success(request, f"Quotation '{quotation.bill_number}' successfully converted to Tax Invoice '{invoice.bill_number}'!")
    return redirect('billing:bill_edit', pk=invoice.pk)


class BillPDFView(LoginRequiredMixin, View):
    """
    Renders high-fidelity downloadable / viewable A4 PDF of the bill using xhtml2pdf.
    Supports copy type parameter: ?copy=ORIGINAL, ?copy=DUPLICATE, ?copy=TRIPLICATE
    and ?download=1 to force attachment download.
    """
    def get(self, request, pk, *args, **kwargs):
        bill = get_object_or_404(Bill.objects.select_related('customer', 'gst_config'), pk=pk)
        company = CompanyProfile.get_instance()
        copy_param = request.GET.get('copy', 'ORIGINAL').upper()

        if 'DUP' in copy_param or 'TRANS' in copy_param:
            copy_type = 'DUPLICATE FOR TRANSPORTER'
        elif 'TRIP' in copy_param or 'OFFICE' in copy_param or 'SHOWROOM' in copy_param:
            copy_type = 'TRIPLICATE FOR SHOWROOM'
        else:
            copy_type = 'ORIGINAL FOR BUYER'

        # Quotation finance estimation
        margin_suggested = (bill.grand_total * Decimal('0.15')).quantize(Decimal('0.01'))
        loan_balance = (bill.grand_total - margin_suggested).quantize(Decimal('0.01'))

        items = list(bill.items.all().order_by('sr_no'))
        total_quantity = sum((item.quantity for item in items), 0)

        context = {
            'bill': bill,
            'items': items,
            'total_quantity': total_quantity,
            'company': company,
            'copy_type': copy_type,
            'hsn_summary': bill.get_hsn_summary(),
            'margin_suggested': margin_suggested,
            'loan_balance': loan_balance,
        }

        html = render_to_string('billing/bill_pdf.html', context)
        result = io.BytesIO()

        pdf = pisa.pisaDocument(io.BytesIO(html.encode("UTF-8")), result, link_callback=link_callback)
        if not pdf.err:
            response = HttpResponse(result.getvalue(), content_type='application/pdf')
            disposition = 'attachment' if request.GET.get('download') == '1' else 'inline'
            filename = f"{bill.bill_number}_{bill.bill_type}.pdf".replace(' ', '_').replace('/', '_')
            response['Content-Disposition'] = f'{disposition}; filename="{filename}"'
            return response

        messages.error(request, "Could not generate PDF document.")
        return redirect('billing:bill_detail', pk=bill.pk)


class TemporaryBillView(LoginRequiredMixin, View):
    """
    Temporary / Scratch Bill Control Panel:
    100% Client-Side (HTML/CSS/JS) interactive bill generator and printer.
    Allows staff to quickly fill in details and print professional estimates, quotations,
    counter cash memos, or delivery slips with zero database writes, zero inventory deduction,
    and zero accounting impact.
    """
    template_name = 'billing/temporary_bill.html'

    def get(self, request, *args, **kwargs):
        company = CompanyProfile.get_instance()
        primary_bank = company.primary_bank
        return render(request, self.template_name, {
            'company': company,
            'primary_bank': primary_bank,
            'page_title': 'Temporary Bill (Print-Only)',
        })


class InventoryDashboardView(LoginRequiredMixin, ListView):
    """
    Central Showroom Inventory Management Dashboard.
    Provides real-time visibility into showroom stock: available units,
    valuation, low-stock warnings, serialized tractors, and inward purchases.
    """
    model = Product
    template_name = 'billing/inventory_dashboard.html'
    context_object_name = 'products'
    paginate_by = 20

    def get_paginate_by(self, queryset):
        page_size = self.request.GET.get('page_size')
        if page_size and page_size.isdigit():
            size = int(page_size)
            if 5 <= size <= 100:
                return size
        return self.paginate_by

    def get_queryset(self):
        qs = Product.objects.filter(is_active=True).order_by('category', 'name')
        category = self.request.GET.get('category', '').strip()
        status = self.request.GET.get('status', '').strip()
        q = self.request.GET.get('q', '').strip()

        if category:
            qs = qs.filter(category=category)
        if q:
            qs = qs.filter(
                Q(name__icontains=q) |
                Q(hsn_code__icontains=q) |
                Q(part_number__icontains=q)
            )

        if status == 'IN_STOCK':
            qs = qs.filter(stock_quantity__gt=0)
        elif status == 'OUT_OF_STOCK':
            qs = qs.filter(stock_quantity=0)
        elif status == 'LOW_STOCK':
            from django.db.models import F
            qs = qs.filter(stock_quantity__gt=0, stock_quantity__lte=F('min_stock_level'))

        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        all_active = Product.objects.filter(is_active=True)

        stats = all_active.aggregate(
            total_qty=Sum('stock_quantity'),
            total_valuation=Sum(
                F('stock_quantity') * Coalesce(
                    'purchase_price',
                    'default_price',
                    Value(Decimal('0.00')),
                    output_field=DecimalField()
                )
            ),
            low_stock_count=Count('id', filter=Q(stock_quantity__gt=0, stock_quantity__lte=F('min_stock_level'))),
            out_of_stock_count=Count('id', filter=Q(stock_quantity=0)),
            in_stock_count=Count('id', filter=Q(stock_quantity__gt=0)),
        )

        total_qty = stats['total_qty'] or 0
        total_valuation = stats['total_valuation'] or Decimal('0.00')
        low_stock_count = stats['low_stock_count'] or 0
        out_of_stock_count = stats['out_of_stock_count'] or 0
        in_stock_count = stats['in_stock_count'] or 0

        # Serialized tractors
        tractors_in_stock = TractorStockUnit.objects.filter(status='IN_STOCK').count()
        serialized_units = TractorStockUnit.objects.select_related('product', 'allocated_bill').all().order_by('-created_at')[:10]

        # Recent Purchases
        recent_purchases = PurchaseBill.objects.prefetch_related('items').order_by('-invoice_date')[:5]

        context.update({
            'current_page_size': self.get_paginate_by(self.get_queryset()),
            'current_category': self.request.GET.get('category', ''),
            'current_status': self.request.GET.get('status', ''),
            'search_query': self.request.GET.get('q', ''),
            'categories': Product.CATEGORY_CHOICES,
            'total_stock_qty': total_qty,
            'total_stock_valuation': total_valuation,
            'low_stock_count': low_stock_count,
            'out_of_stock_count': out_of_stock_count,
            'in_stock_count': in_stock_count,
            'tractors_in_stock': tractors_in_stock,
            'serialized_units': serialized_units,
            'recent_purchases': recent_purchases,
        })
        return context


class PurchaseBillListView(LoginRequiredMixin, ListView):
    """
    List of Inward Stock / Manufacturer Purchase Invoices.
    """
    model = PurchaseBill
    template_name = 'billing/purchase_list.html'
    context_object_name = 'purchases'
    paginate_by = 15

    def get_queryset(self):
        qs = PurchaseBill.objects.prefetch_related('items__product').all().order_by('-invoice_date', '-created_at')
        q = self.request.GET.get('q', '').strip()
        p_type = self.request.GET.get('type', '').strip()
        if q:
            qs = qs.filter(
                Q(invoice_number__icontains=q) |
                Q(supplier_name__icontains=q) |
                Q(supplier_gstin__icontains=q)
            )
        if p_type:
            qs = qs.filter(purchase_type=p_type)
        return qs

    def get_paginate_by(self, queryset):
        page_size = self.request.GET.get('page_size')
        if page_size and page_size.isdigit():
            size = int(page_size)
            if 5 <= size <= 100:
                return size
        return self.paginate_by

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        context['current_page_size'] = self.get_paginate_by(self.get_queryset())
        context['search_query'] = self.request.GET.get('q', '')
        context['current_type'] = self.request.GET.get('type', '')
        context['types'] = PurchaseBill.PURCHASE_TYPE_CHOICES

        kpis = PurchaseBill.objects.aggregate(
            total_inward_amount=Sum('total_amount'),
            total_inward_bills=Count('id'),
            tractor_count=Count('id', filter=Q(purchase_type='TRACTOR')),
            tractor_sum=Sum('total_amount', filter=Q(purchase_type='TRACTOR')),
            implement_count=Count('id', filter=Q(purchase_type='IMPLEMENT')),
            implement_sum=Sum('total_amount', filter=Q(purchase_type='IMPLEMENT')),
            spare_count=Count('id', filter=Q(purchase_type='SPARE')),
            spare_sum=Sum('total_amount', filter=Q(purchase_type='SPARE')),
        )

        context['total_inward_amount'] = kpis['total_inward_amount'] or Decimal('0.00')
        context['total_inward_bills'] = kpis['total_inward_bills'] or 0
        context['tractor_bills_count'] = kpis['tractor_count'] or 0
        context['tractor_bills_sum'] = kpis['tractor_sum'] or Decimal('0.00')
        context['implement_bills_count'] = kpis['implement_count'] or 0
        context['implement_bills_sum'] = kpis['implement_sum'] or Decimal('0.00')
        context['spare_bills_count'] = kpis['spare_count'] or 0
        context['spare_bills_sum'] = kpis['spare_sum'] or Decimal('0.00')

        return context


class PurchaseBillCreateView(LoginRequiredMixin, View):
    """
    Record an Inward Stock / Purchase Bill with interactive line items.
    Automatically increments showroom inventory stock upon saving.
    """
    template_name = 'billing/purchase_form.html'

    def get(self, request, *args, **kwargs):
        form = PurchaseBillForm()
        formset = PurchaseBillItemFormSet()
        products_json = json.dumps([
            {
                'id': p.id,
                'name': p.name,
                'category': p.category,
                'category_display': p.get_category_display(),
                'hsn': p.hsn_code or '',
                'cost': float(p.purchase_price or 0),
                'stock': p.stock_quantity,
            }
            for p in Product.objects.filter(is_active=True).order_by('category', 'name')
        ])
        return render(request, self.template_name, {
            'form': form,
            'formset': formset,
            'products_json': products_json,
            'page_title': 'Record Inward Stock (Purchase Invoice)',
        })

    def post(self, request, *args, **kwargs):
        form = PurchaseBillForm(request.POST)
        formset = PurchaseBillItemFormSet(request.POST)

        if form.is_valid() and formset.is_valid():
            with transaction.atomic():
                purchase = form.save(commit=False)
                purchase.created_by = request.user
                purchase.save()

                formset.instance = purchase
                formset.save()

                # Calculate total amount
                total_amt = sum((item.amount for item in purchase.items.all()), Decimal('0.00'))
                purchase.total_amount = total_amt
                purchase.save(update_fields=['total_amount'])

                # Apply stock inward to showroom inventory
                apply_purchase_stock_inward(purchase)

            messages.success(
                request,
                f"Inward Purchase Invoice '{purchase.invoice_number}' from '{purchase.supplier_name}' "
                f"recorded successfully! Showroom stock updated."
            )
            return redirect('billing:purchase_detail', pk=purchase.pk)

        products_json = json.dumps([
            {
                'id': p.id,
                'name': p.name,
                'category': p.category,
                'category_display': p.get_category_display(),
                'hsn': p.hsn_code or '',
                'cost': float(p.purchase_price or 0),
                'stock': p.stock_quantity,
            }
            for p in Product.objects.filter(is_active=True).order_by('category', 'name')
        ])
        messages.error(request, "Please correct the highlighted errors below before saving purchase invoice.")
        return render(request, self.template_name, {
            'form': form,
            'formset': formset,
            'products_json': products_json,
            'page_title': 'Record Inward Stock (Purchase Invoice)',
        })


class PurchaseBillDetailView(LoginRequiredMixin, DetailView):
    """
    Displays the inward purchase bill details, supplier particulars, and items received.
    """
    model = PurchaseBill
    template_name = 'billing/purchase_detail.html'
    context_object_name = 'purchase'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['items'] = self.object.items.select_related('product').all()
        return context


class ExportInventoryExcelView(LoginRequiredMixin, View):
    """
    Exports showroom inventory to a professionally formatted Excel spreadsheet.
    """
    def get(self, request, *args, **kwargs):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        company = CompanyProfile.get_instance()
        products = Product.objects.filter(is_active=True).order_by('category', 'name')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Showroom Inventory"

        title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
        subtitle_font = Font(name="Calibri", size=10, italic=True, color="475569")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        total_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )
        double_bottom_border = Border(
            top=Side(style='thin', color='0F172A'),
            bottom=Side(style='double', color='0F172A')
        )

        ws.append([company.company_name or "Tractor Showroom Management ERP"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Showroom Stock & Inventory Valuation Report — Generated: {timezone.now().strftime('%d/%m/%Y %H:%M')}"])
        ws.cell(row=2, column=1).font = subtitle_font
        ws.append([])

        headers = [
            "Sr", "Item / Model Description", "Category", "SKU / Part No",
            "HSN Code", "Showroom Stock Qty", "Procurement Cost (₹)", "Selling Price (₹)",
            "Total Valuation (₹)", "Stock Status"
        ]
        ws.append(headers)
        header_row = 4
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        start_row = 5
        curr_row = start_row
        for i, p in enumerate(products, start=1):
            status_display = "IN STOCK" if p.stock_quantity > p.min_stock_level else ("LOW STOCK" if p.stock_quantity > 0 else "OUT OF STOCK")
            ws.append([
                i,
                p.name,
                p.get_category_display(),
                p.part_number or "",
                p.hsn_code or "",
                p.stock_quantity,
                float(p.purchase_price or 0),
                float(p.default_price or 0),
                float(p.total_stock_value or 0),
                status_display,
            ])
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_row, column=c)
                cell.border = thin_border
                if c in [7, 8, 9]:
                    cell.number_format = '₹#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif c in [1, 5, 6, 10]:
                    cell.alignment = Alignment(horizontal="center")
            curr_row += 1

        if curr_row > start_row:
            end_row = curr_row - 1
            ws.append([
                "TOTAL INVENTORY VALUATION", "", "", "", "",
                f"=SUM(F{start_row}:F{end_row})",
                "", "",
                f"=SUM(I{start_row}:I{end_row})",
                ""
            ])
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_row, column=c)
                cell.font = total_font
                cell.fill = total_fill
                cell.border = double_bottom_border
                if c in [7, 8, 9]:
                    cell.number_format = '₹#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif c == 6:
                    cell.alignment = Alignment(horizontal="center")

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"Inventory_Report_{timezone.now().strftime('%Y%m%d')}.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class DownloadSampleInventoryView(LoginRequiredMixin, View):
    """
    Downloads a sample Excel template for bulk importing products & stock.
    """
    def get(self, request, *args, **kwargs):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sample Template"

        headers = [
            "Product Name", "Category (TRACTOR/IMPLEMENT/TROLLEY/SPARE/ACCESSORY)",
            "Part Number", "HSN Code", "Selling Price", "Cost Price", "Stock Quantity", "Min Stock Level"
        ]
        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        samples = [
            ["Captain 283 DI 4WD", "TRACTOR", "CPT-283-4WD", "8701", 595000.00, 520000.00, 2, 1],
            ["Cultivator 9 Tyne Rigid", "IMPLEMENT", "CLT-09-RGD", "8432", 38000.00, 31000.00, 5, 2],
            ["Engine Oil 15W40 CI-4 (7.5L)", "SPARE", "LUB-15W40-75", "2710", 2850.00, 2250.00, 25, 5],
            ["Hydraulic Trolley 3 Ton", "TROLLEY", "TRL-3T-SNGL", "8716", 125000.00, 105000.00, 3, 1],
            ["Tractor Front Tyre 6.00-16", "SPARE", "TYR-600-16", "4011", 5800.00, 4600.00, 10, 4],
        ]
        for row in samples:
            ws.append(row)

        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 24

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = 'attachment; filename="Inventory_Import_Template.xlsx"'
        return response


class ImportInventoryExcelView(LoginRequiredMixin, View):
    """
    Bulk import/update inventory products and stock quantities from Excel.
    """
    template_name = 'billing/inventory_import.html'

    def get(self, request, *args, **kwargs):
        return render(request, self.template_name, {
            'page_title': 'Import Inventory & Stock',
        })

    def post(self, request, *args, **kwargs):
        import openpyxl
        file = request.FILES.get('file')
        if not file:
            messages.error(request, "Please choose an Excel file (.xlsx) to upload.")
            return redirect('billing:inventory_import')

        try:
            wb = openpyxl.load_workbook(file, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                messages.error(request, "The uploaded Excel file has no data rows.")
                return redirect('billing:inventory_import')

            created_count = 0
            updated_count = 0

            with transaction.atomic():
                for row_idx, row in enumerate(rows[1:], start=2):
                    if not row or not row[0]:
                        continue
                    name = str(row[0]).strip()
                    cat = str(row[1]).strip().upper() if len(row) > 1 and row[1] else 'SPARE'
                    valid_cats = [c[0] for c in Product.CATEGORY_CHOICES]
                    if cat not in valid_cats:
                        cat = 'SPARE'

                    part_no = str(row[2]).strip() if len(row) > 2 and row[2] else ''
                    hsn = str(row[3]).strip() if len(row) > 3 and row[3] else ('8701' if cat == 'TRACTOR' else '8432')
                    import re
                    def parse_dec(val):
                        if val is None:
                            return Decimal('0.00')
                        cleaned = re.sub(r'[^\d.]', '', str(val))
                        try:
                            return Decimal(cleaned) if cleaned else Decimal('0.00')
                        except Exception:
                            return Decimal('0.00')

                    def parse_int(val, default=0):
                        if val is None:
                            return default
                        try:
                            cleaned = re.sub(r'[^\d.]', '', str(val))
                            return int(float(cleaned)) if cleaned else default
                        except Exception:
                            return default

                    price = parse_dec(row[4]) if len(row) > 4 else Decimal('0.00')
                    cost = parse_dec(row[5]) if len(row) > 5 else Decimal('0.00')
                    qty = parse_int(row[6], 0) if len(row) > 6 else 0
                    min_stock = parse_int(row[7], 2) if len(row) > 7 else 2

                    prod = Product.objects.filter(name__iexact=name).first()
                    if not prod:
                        prod = Product.objects.create(
                            name=name,
                            category=cat,
                            part_number=part_no,
                            hsn_code=hsn,
                            default_price=price,
                            purchase_price=cost,
                            stock_quantity=qty,
                            min_stock_level=min_stock,
                            is_active=True,
                        )
                        created_count += 1
                    else:
                        prod.stock_quantity = qty
                        if price > Decimal('0.00'):
                            prod.default_price = price
                        if cost > Decimal('0.00'):
                            prod.purchase_price = cost
                        if part_no:
                            prod.part_number = part_no
                        prod.is_active = True
                        prod.save()
                        updated_count += 1

            messages.success(
                request,
                f"Inventory successfully imported! {created_count} new product(s) added, "
                f"{updated_count} existing product(s) updated."
            )
            return redirect('billing:inventory_dashboard')

        except Exception as e:
            messages.error(request, f"Failed to import Excel file: {str(e)}")
            return redirect('billing:inventory_import')


class ProductListView(LoginRequiredMixin, ListView):
    """Catalog of tractor models, implements, attachments, and spare parts."""
    model = Product
    template_name = 'billing/product_list.html'
    context_object_name = 'products'
    paginate_by = 15

    def get_paginate_by(self, queryset):
        page_size = self.request.GET.get('page_size')
        if page_size and page_size.isdigit():
            size = int(page_size)
            if 5 <= size <= 100:
                return size
        return self.paginate_by

    def get_queryset(self):
        qs = Product.objects.all().order_by('category', 'name')
        category = self.request.GET.get('category', '').strip()
        q = self.request.GET.get('q', '').strip()

        if category:
            qs = qs.filter(category=category)
        if q:
            qs = qs.filter(
                Q(name__icontains=q) |
                Q(hsn_code__icontains=q) |
                Q(part_number__icontains=q)
            )
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['current_page_size'] = self.get_paginate_by(self.get_queryset())
        context['current_category'] = self.request.GET.get('category', '')
        context['search_query'] = self.request.GET.get('q', '')
        context['categories'] = Product.CATEGORY_CHOICES
        return context


class ProductCreateView(LoginRequiredMixin, View):
    """Add a new tractor model, implement, or spare part to catalog."""
    template_name = 'billing/product_form.html'

    def get(self, request, *args, **kwargs):
        form = ProductForm()
        return render(request, self.template_name, {
            'form': form,
            'is_edit': False,
            'page_title': 'Add Product / Model / Spare Part',
        })

    def post(self, request, *args, **kwargs):
        form = ProductForm(request.POST)
        if form.is_valid():
            product = form.save()
            messages.success(request, f"Product '{product.name}' created successfully!")
            return redirect('billing:product_list')
        return render(request, self.template_name, {
            'form': form,
            'is_edit': False,
            'page_title': 'Add Product / Model / Spare Part',
        })


class ProductUpdateView(LoginRequiredMixin, View):
    """Edit an existing product model in catalog."""
    template_name = 'billing/product_form.html'

    def get(self, request, pk, *args, **kwargs):
        product = get_object_or_404(Product, pk=pk)
        form = ProductForm(instance=product)
        return render(request, self.template_name, {
            'form': form,
            'is_edit': True,
            'product': product,
            'page_title': f"Edit Product - {product.name}",
        })

    def post(self, request, pk, *args, **kwargs):
        product = get_object_or_404(Product, pk=pk)
        form = ProductForm(request.POST, instance=product)
        if form.is_valid():
            product = form.save()
            messages.success(request, f"Product '{product.name}' updated successfully!")
            return redirect('billing:product_list')
        return render(request, self.template_name, {
            'form': form,
            'is_edit': True,
            'product': product,
            'page_title': f"Edit Product - {product.name}",
        })


class TractorStockUnitCreateView(LoginRequiredMixin, View):
    """Add a new serialized tractor unit to showroom inventory."""
    template_name = 'billing/stock_unit_form.html'

    def get(self, request, *args, **kwargs):
        form = TractorStockUnitForm()
        return render(request, self.template_name, {
            'form': form,
            'is_edit': False,
            'page_title': 'Add Serialized Tractor Unit (Stock In)',
        })

    def post(self, request, *args, **kwargs):
        form = TractorStockUnitForm(request.POST)
        if form.is_valid():
            unit = form.save()
            # Update product stock count
            prod = unit.product
            prod.stock_quantity = TractorStockUnit.objects.filter(product=prod, status='IN_STOCK').count()
            prod.save(update_fields=['stock_quantity'])

            messages.success(request, f"Tractor Unit with Chassis '{unit.chassis_number}' added to inventory successfully!")
            return redirect('billing:inventory_dashboard')
        return render(request, self.template_name, {
            'form': form,
            'is_edit': False,
            'page_title': 'Add Serialized Tractor Unit (Stock In)',
        })


class TractorStockUnitUpdateView(LoginRequiredMixin, View):
    """Edit a serialized tractor unit in inventory."""
    template_name = 'billing/stock_unit_form.html'

    def get(self, request, pk, *args, **kwargs):
        unit = get_object_or_404(TractorStockUnit, pk=pk)
        form = TractorStockUnitForm(instance=unit)
        return render(request, self.template_name, {
            'form': form,
            'is_edit': True,
            'unit': unit,
            'page_title': f"Edit Tractor Unit - Chassis {unit.chassis_number}",
        })

    def post(self, request, pk, *args, **kwargs):
        unit = get_object_or_404(TractorStockUnit, pk=pk)
        form = TractorStockUnitForm(request.POST, instance=unit)
        if form.is_valid():
            unit = form.save()
            # Update product stock count
            prod = unit.product
            prod.stock_quantity = TractorStockUnit.objects.filter(product=prod, status='IN_STOCK').count()
            prod.save(update_fields=['stock_quantity'])

            messages.success(request, f"Tractor Unit '{unit.chassis_number}' updated successfully!")
            return redirect('billing:inventory_dashboard')
        return render(request, self.template_name, {
            'form': form,
            'is_edit': True,
            'unit': unit,
            'page_title': f"Edit Tractor Unit - Chassis {unit.chassis_number}",
        })


class TractorStockUnitDeleteView(LoginRequiredMixin, View):
    """Delete a tractor unit from inventory."""
    def post(self, request, pk, *args, **kwargs):
        unit = get_object_or_404(TractorStockUnit, pk=pk)
        chassis = unit.chassis_number
        prod = unit.product
        unit.delete()
        if prod:
            prod.stock_quantity = TractorStockUnit.objects.filter(product=prod, status='IN_STOCK').count()
            prod.save(update_fields=['stock_quantity'])
        messages.success(request, f"Tractor unit with Chassis '{chassis}' deleted.")
        return redirect('billing:inventory_dashboard')

