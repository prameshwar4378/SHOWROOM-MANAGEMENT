from decimal import Decimal
from django.shortcuts import render
from django.views import View
from django.views.generic import TemplateView
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import Sum, Count, Q, Subquery, OuterRef, Value, DecimalField, F
from django.db.models.functions import Coalesce
from django.utils import timezone
from datetime import timedelta


class DashboardView(LoginRequiredMixin, TemplateView):
    """Main dashboard showing key business metrics and recent activity."""
    template_name = 'reports/dashboard.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = timezone.now().date()
        month_start = today.replace(day=1)

        try:
            from customers.models import Customer
            from billing.models import Bill
            from payments.models import Payment

            # Customer stats
            context['total_customers'] = Customer.objects.filter(is_active=True).count()
            context['new_customers_month'] = Customer.objects.filter(
                is_active=True, created_at__date__gte=month_start
            ).count()

            # Bill stats (combined into single aggregate query)
            bill_aggregates = Bill.objects.exclude(status='CANCELLED').aggregate(
                total_bills=Count('id'),
                bills_this_month=Count('id', filter=Q(bill_date__gte=month_start)),
                total_billed_amount=Sum('grand_total', filter=Q(bill_type='TAX_INVOICE')),
                billed_this_month=Sum('grand_total', filter=Q(bill_type='TAX_INVOICE', bill_date__gte=month_start)),
                quotation_count=Count('id', filter=Q(bill_type='QUOTATION')),
                tax_invoice_count=Count('id', filter=Q(bill_type='TAX_INVOICE')),
                margin_money_count=Count('id', filter=Q(bill_type='MARGIN_MONEY')),
                delivery_challan_count=Count('id', filter=Q(bill_type='DELIVERY_CHALLAN')),
            )
            context['total_bills'] = bill_aggregates['total_bills'] or 0
            context['bills_this_month'] = bill_aggregates['bills_this_month'] or 0
            context['total_billed_amount'] = bill_aggregates['total_billed_amount'] or Decimal('0.00')
            context['billed_this_month'] = bill_aggregates['billed_this_month'] or Decimal('0.00')
            context['quotation_count'] = bill_aggregates['quotation_count'] or 0
            context['tax_invoice_count'] = bill_aggregates['tax_invoice_count'] or 0
            context['margin_money_count'] = bill_aggregates['margin_money_count'] or 0
            context['delivery_challan_count'] = bill_aggregates['delivery_challan_count'] or 0

            # Payment stats (combined into single aggregate query)
            payment_aggregates = Payment.objects.aggregate(
                total_collected=Sum('amount'),
                collected_this_month=Sum('amount', filter=Q(payment_date__gte=month_start)),
                collected_today=Sum('amount', filter=Q(payment_date=today)),
            )
            context['total_collected'] = payment_aggregates['total_collected'] or Decimal('0.00')
            context['collected_this_month'] = payment_aggregates['collected_this_month'] or Decimal('0.00')
            context['collected_today'] = payment_aggregates['collected_today'] or Decimal('0.00')

            # Recent activity
            context['recent_bills'] = Bill.objects.select_related(
                'customer'
            ).order_by('-created_at')[:5]
            context['recent_payments'] = Payment.objects.select_related(
                'customer'
            ).order_by('-created_at')[:5]

            # Outstanding & Customers with pending payments - optimized via Subquery annotations
            b_sq = Bill.objects.filter(
                customer=OuterRef('pk'),
                bill_type='TAX_INVOICE'
            ).exclude(status='CANCELLED').values('customer').annotate(tot=Sum('grand_total')).values('tot')
            p_sq = Payment.objects.filter(
                customer=OuterRef('pk')
            ).values('customer').annotate(tot=Sum('amount')).values('tot')

            pending_qs = (
                Customer.objects.filter(is_active=True)
                .annotate(
                    db_billed=Coalesce(Subquery(b_sq), Value(Decimal('0.00')), output_field=DecimalField()),
                    db_paid=Coalesce(Subquery(p_sq), Value(Decimal('0.00')), output_field=DecimalField()),
                )
                .annotate(db_due=F('db_billed') - F('db_paid'))
                .filter(db_due__gt=Decimal('0.00'))
            )
            total_out = pending_qs.aggregate(tot=Sum('db_due'))['tot'] or Decimal('0.00')
            context['total_outstanding'] = total_out
            context['pending_customers'] = list(pending_qs.order_by('-db_due')[:5])

            # Inventory & machinery stock stats
            from billing.models import Product, TractorStockUnit
            total_tractors = TractorStockUnit.objects.count()
            available_tractors = TractorStockUnit.objects.filter(status='AVAILABLE').count()
            allocated_tractors = TractorStockUnit.objects.filter(status='ALLOCATED').count()
            total_products = Product.objects.count()
            low_stock_products = Product.objects.filter(stock_quantity__lte=2, stock_quantity__gt=0).count()
            out_of_stock_products = Product.objects.filter(stock_quantity=0).count()

            context['inventory_stats'] = {
                'total_tractors': total_tractors,
                'available_tractors': available_tractors,
                'allocated_tractors': allocated_tractors,
                'total_products': total_products,
                'low_stock_products': low_stock_products,
                'out_of_stock_products': out_of_stock_products,
            }

            # Showroom operating expenses stats
            from expenses.models import Expense
            exp_aggs = Expense.objects.aggregate(
                total_expenses=Sum('amount'),
                expenses_this_month=Sum('amount', filter=Q(expense_date__gte=month_start)),
                expenses_today=Sum('amount', filter=Q(expense_date=today)),
                expenses_count=Count('id'),
            )
            total_expenses = exp_aggs['total_expenses'] or Decimal('0.00')
            expenses_this_month = exp_aggs['expenses_this_month'] or Decimal('0.00')
            expenses_today = exp_aggs['expenses_today'] or Decimal('0.00')

            context['total_expenses'] = total_expenses
            context['expenses_this_month'] = expenses_this_month
            context['expenses_today'] = expenses_today
            context['net_cash_flow'] = context['total_collected'] - total_expenses

            # Collection efficiency
            billed_tot = context['total_billed_amount']
            col_tot = context['total_collected']
            context['collection_rate'] = round(float((col_tot / billed_tot) * 100), 1) if billed_tot > 0 else 0.0

            # 14-Day Daily Trends (for smooth Spline Chart)
            from datetime import date
            daily_labels = []
            daily_billed = []
            daily_collected = []
            for i in range(13, -1, -1):
                d = today - timedelta(days=i)
                daily_labels.append(d.strftime('%d %b'))
                daily_billed.append(float(Bill.objects.filter(bill_type='TAX_INVOICE', bill_date=d).exclude(status='CANCELLED').aggregate(s=Sum('grand_total'))['s'] or 0))
                daily_collected.append(float(Payment.objects.filter(payment_date=d).aggregate(s=Sum('amount'))['s'] or 0))

            # 6-Month Historical Trends
            monthly_labels = []
            monthly_billed = []
            monthly_collected = []
            monthly_expenses = []
            for i in range(5, -1, -1):
                m = today.month - i
                y = today.year
                while m <= 0:
                    m += 12
                    y -= 1
                sd = date(y, m, 1)
                ed = date(y + 1, 1, 1) if m == 12 else date(y, m + 1, 1)
                monthly_labels.append(sd.strftime('%b %Y'))
                monthly_billed.append(float(Bill.objects.filter(bill_type='TAX_INVOICE', bill_date__gte=sd, bill_date__lt=ed).exclude(status='CANCELLED').aggregate(s=Sum('grand_total'))['s'] or 0))
                monthly_collected.append(float(Payment.objects.filter(payment_date__gte=sd, payment_date__lt=ed).aggregate(s=Sum('amount'))['s'] or 0))
                monthly_expenses.append(float(Expense.objects.filter(expense_date__gte=sd, expense_date__lt=ed).aggregate(s=Sum('amount'))['s'] or 0))

            # Payment modes distribution
            pm_map = dict(Payment.PAYMENT_MODE_CHOICES)
            pm_qs = Payment.objects.values('payment_mode').annotate(total=Sum('amount'), count=Count('id')).order_by('-total')
            pm_list = []
            total_pm_amount = float(sum(x['total'] or 0 for x in pm_qs)) or 1.0
            palette = ['#06b6d4', '#10b981', '#f59e0b', '#6366f1', '#ec4899', '#8b5cf6']
            for idx, pm in enumerate(pm_qs):
                t = float(pm['total'] or 0)
                pm_list.append({
                    'code': pm['payment_mode'],
                    'label': str(pm_map.get(pm['payment_mode'], pm['payment_mode'])),
                    'total': t,
                    'count': pm['count'],
                    'percent': round((t / total_pm_amount) * 100, 1),
                    'color': palette[idx % len(palette)],
                })
            context['payment_modes_breakdown'] = pm_list

            import json
            chart_payload = {
                'daily_labels': daily_labels,
                'daily_billed': daily_billed,
                'daily_collected': daily_collected,
                'monthly_labels': monthly_labels,
                'monthly_billed': monthly_billed,
                'monthly_collected': monthly_collected,
                'monthly_expenses': monthly_expenses,
                'doc_labels': ['Tax Invoices', 'Quotations', 'Margin Money', 'Delivery Challans'],
                'doc_counts': [
                    context['tax_invoice_count'],
                    context['quotation_count'],
                    context['margin_money_count'],
                    context['delivery_challan_count'],
                ],
                'doc_colors': ['#10b981', '#3b82f6', '#8b5cf6', '#f97316'],
                'pm_labels': [x['label'] for x in pm_list],
                'pm_totals': [x['total'] for x in pm_list],
                'pm_colors': [x['color'] for x in pm_list],
            }
            context['chart_data_json'] = json.dumps(chart_payload)

        except Exception as e:
            # Apps might not be migrated yet or initial setup
            import logging
            logging.getLogger(__name__).warning(f"Dashboard context aggregation notice: {e}")
            context.update({
                'total_customers': 0, 'new_customers_month': 0,
                'total_bills': 0, 'bills_this_month': 0,
                'total_billed_amount': 0, 'billed_this_month': 0,
                'quotation_count': 0, 'tax_invoice_count': 0,
                'margin_money_count': 0, 'delivery_challan_count': 0,
                'total_collected': 0, 'collected_this_month': 0,
                'collected_today': 0, 'total_outstanding': 0,
                'recent_bills': [], 'recent_payments': [],
                'pending_customers': [], 'chart_data_json': '{}',
                'collection_rate': 0.0, 'total_expenses': 0,
                'expenses_this_month': 0, 'expenses_today': 0,
                'net_cash_flow': 0,
                'inventory_stats': {'total_tractors': 0, 'available_tractors': 0, 'allocated_tractors': 0, 'total_products': 0, 'low_stock_products': 0, 'out_of_stock_products': 0},
                'payment_modes_breakdown': [],
            })

        return context


class ExportSalesExcelView(LoginRequiredMixin, View):
    """
    Exports showroom sales & bill records to a professionally formatted Excel spreadsheet.
    Supports GSTR-1 and tax audit reconciliation.
    """
    def get(self, request, *args, **kwargs):
        import io
        from decimal import Decimal
        from django.http import HttpResponse
        from django.views import View
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from billing.models import Bill
        from settings_app.models import CompanyProfile

        company = CompanyProfile.get_instance()
        bills = Bill.objects.select_related('customer').exclude(status='CANCELLED').order_by('-bill_date', '-created_at')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales Register"

        title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
        subtitle_font = Font(name="Calibri", size=10, italic=True, color="4B5563")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        total_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
        total_font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
        thin_border = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB')
        )
        double_bottom_border = Border(
            top=Side(style='thin', color='1E40AF'),
            bottom=Side(style='double', color='1E40AF')
        )

        ws.append([company.company_name or "Showroom Management ERP"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Sales Register & Bill Records Report — Generated: {timezone.now().strftime('%d/%m/%Y %H:%M')}"])
        ws.cell(row=2, column=1).font = subtitle_font
        ws.append([])

        headers = [
            "Sr", "Bill No", "Document Type", "Date", "Customer Name", "Mobile",
            "Village / City", "GST Bill?", "Subtotal (₹)", "Discount (₹)",
            "Taxable (₹)", "CGST (₹)", "SGST (₹)", "IGST (₹)", "Total Tax (₹)",
            "Round-Off (₹)", "Grand Total (₹)", "Status"
        ]
        ws.append(headers)
        header_row = 4
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        start_row = 5
        curr_row = start_row
        for i, b in enumerate(bills, start=1):
            ws.append([
                i,
                b.bill_number,
                b.get_bill_type_display(),
                b.bill_date.strftime('%d/%m/%Y') if b.bill_date else "",
                b.customer.full_name if b.customer else "",
                b.customer.mobile_primary if b.customer else "",
                b.customer.village if b.customer else "",
                "YES" if b.is_gst_bill else "NO",
                float(b.subtotal or 0),
                float(b.discount_amount or 0),
                float(b.taxable_amount or 0),
                float(b.cgst_amount or 0),
                float(b.sgst_amount or 0),
                float(b.igst_amount or 0),
                float(b.total_tax or 0),
                float(b.round_off_amount or 0),
                float(b.grand_total or 0),
                b.status,
            ])
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_row, column=c)
                cell.border = thin_border
                if c in [9, 10, 11, 12, 13, 14, 15, 16, 17]:
                    cell.number_format = '₹#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif c in [1, 4, 8, 18]:
                    cell.alignment = Alignment(horizontal="center")
            curr_row += 1

        if curr_row > start_row:
            end_row = curr_row - 1
            ws.append([
                "TOTAL", "", "", "", "", "", "", "",
                f"=SUM(I{start_row}:I{end_row})",
                f"=SUM(J{start_row}:J{end_row})",
                f"=SUM(K{start_row}:K{end_row})",
                f"=SUM(L{start_row}:L{end_row})",
                f"=SUM(M{start_row}:M{end_row})",
                f"=SUM(N{start_row}:N{end_row})",
                f"=SUM(O{start_row}:O{end_row})",
                f"=SUM(P{start_row}:P{end_row})",
                f"=SUM(Q{start_row}:Q{end_row})",
                ""
            ])
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_row, column=c)
                cell.font = total_font
                cell.fill = total_fill
                cell.border = double_bottom_border
                if c in [9, 10, 11, 12, 13, 14, 15, 16, 17]:
                    cell.number_format = '₹#,##0.00'
                    cell.alignment = Alignment(horizontal="right")

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"Sales_Register_{timezone.now().strftime('%Y%m%d')}.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class ExportPendingDuesExcelView(LoginRequiredMixin, View):
    """
    Exports customer pending balances & outstanding dues to an Excel spreadsheet.
    """
    def get(self, request, *args, **kwargs):
        import io
        from decimal import Decimal
        from django.http import HttpResponse
        from django.views import View
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from customers.models import Customer
        from settings_app.models import CompanyProfile

        company = CompanyProfile.get_instance()
        from billing.models import Bill
        from payments.models import Payment

        b_sq = Bill.objects.filter(
            customer=OuterRef('pk'),
            bill_type='TAX_INVOICE'
        ).exclude(status='CANCELLED').values('customer').annotate(tot=Sum('grand_total')).values('tot')
        p_sq = Payment.objects.filter(
            customer=OuterRef('pk')
        ).values('customer').annotate(tot=Sum('amount')).values('tot')

        pending = list(
            Customer.objects.filter(is_active=True)
            .annotate(
                db_billed=Coalesce(Subquery(b_sq), Value(Decimal('0.00')), output_field=DecimalField()),
                db_paid=Coalesce(Subquery(p_sq), Value(Decimal('0.00')), output_field=DecimalField()),
            )
            .annotate(db_due=F('db_billed') - F('db_paid'))
            .filter(db_due__gt=Decimal('0.00'))
            .order_by('-db_due')
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Customer Pending Dues"

        title_font = Font(name="Calibri", size=16, bold=True, color="991B1B")
        subtitle_font = Font(name="Calibri", size=10, italic=True, color="4B5563")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="B91C1C", end_color="B91C1C", fill_type="solid")
        total_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        total_font = Font(name="Calibri", size=11, bold=True, color="991B1B")
        thin_border = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB')
        )
        double_bottom_border = Border(
            top=Side(style='thin', color='B91C1C'),
            bottom=Side(style='double', color='B91C1C')
        )

        ws.append([company.company_name or "Showroom Management ERP"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Customer Outstanding Dues Report — Generated: {timezone.now().strftime('%d/%m/%Y %H:%M')}"])
        ws.cell(row=2, column=1).font = subtitle_font
        ws.append([])

        headers = [
            "Sr", "Customer Name", "Primary Mobile", "Secondary Mobile",
            "Village", "Taluka", "District", "Bank", "Branch",
            "Total Billed (₹)", "Total Paid (₹)", "Net Outstanding Due (₹)", "Unallocated Credit (₹)"
        ]
        ws.append(headers)
        header_row = 4
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        start_row = 5
        curr_row = start_row
        for i, c in enumerate(pending, start=1):
            ws.append([
                i,
                c.full_name,
                c.mobile_primary,
                c.mobile_secondary or "",
                c.village or "",
                c.taluka or "",
                c.district or "",
                c.bank_name or "",
                c.bank_branch or "",
                float(c.db_billed or 0),
                float(c.db_paid or 0),
                float(c.db_due or 0),
                float(c.unallocated_balance or 0),
            ])
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.border = thin_border
                if col_idx in [10, 11, 12, 13]:
                    cell.number_format = '₹#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx in [1, 3, 4]:
                    cell.alignment = Alignment(horizontal="center")
            curr_row += 1

        if curr_row > start_row:
            end_row = curr_row - 1
            ws.append([
                "TOTAL OUTSTANDING", "", "", "", "", "", "", "", "",
                f"=SUM(J{start_row}:J{end_row})",
                f"=SUM(K{start_row}:K{end_row})",
                f"=SUM(L{start_row}:L{end_row})",
                f"=SUM(M{start_row}:M{end_row})",
            ])
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = total_font
                cell.fill = total_fill
                cell.border = double_bottom_border
                if col_idx in [10, 11, 12, 13]:
                    cell.number_format = '₹#,##0.00'
                    cell.alignment = Alignment(horizontal="right")

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"Outstanding_Dues_{timezone.now().strftime('%Y%m%d')}.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
